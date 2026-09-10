"""Bulk SSI XLSX Fundamental Financial Statement Ingestion Engine for QPort.

Parses SSI-exported financial statement workbooks, computes normalized semantic financial
payload hashes, preserves raw SSI evidence and lineage, and populates PostgreSQL Finance DB canonical facts.
"""

from __future__ import annotations

import argparse
from concurrent.futures import ThreadPoolExecutor, as_completed
import csv
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import re
import sys
from typing import Any

import openpyxl

from portfolio.finance_catalog import FINANCE_SCHEMA, _schema_connection
from portfolio.financial_data.models import (
    CanonicalFact,
    ConsolidationScope,
    FactIdentityKey,
    PeriodType,
    QualityStatus,
    StatementType,
)


STATEMENT_TYPE_MAP = {
    "BALANCE_SHEET": "BALANCE_SHEET",
    "BALANCESHEET": "BALANCE_SHEET",
    "INCOME_STATEMENT": "INCOME_STATEMENT",
    "INCOMESTATEMENT": "INCOME_STATEMENT",
    "CASH_FLOW": "CASH_FLOW",
    "CASHFLOW": "CASH_FLOW",
    "CASH_FLOW_STATEMENT": "CASH_FLOW",
}

CANONICAL_LINE_MAPPINGS: dict[str, dict[str, str]] = {
    "INCOME_STATEMENT": {
        "tong doanh thu": "IS.REVENUE.TOTAL",
        "doanh thu thuan": "IS.REVENUE.TOTAL",
        "doanh thu hoat dong": "IS.REVENUE.TOTAL",
        "loi nhuan gop": "IS.PROFIT.GROSS",
        "loi nhuan tu hoat dong kinh doanh": "IS.PROFIT.OPERATING",
        "loi nhuan thuan tu hoat dong kinh doanh": "IS.PROFIT.OPERATING",
        "loi nhuan sau thue": "IS.PROFIT.NET",
        "loi nhuan sau thue tndn": "IS.PROFIT.NET",
        "loi nhuan thuan": "IS.PROFIT.NET",
        "loi nhuan truoc thue": "IS.PROFIT.BEFORE_TAX",
        "co phieu dang luu hanh": "IS.SHARES.OUTSTANDING",
        "so luong co phieu luu hanh": "IS.SHARES.OUTSTANDING",
        "chi phi lai vay": "IS.EXPENSE.INTEREST",
        "chi phi quan ly doanh nghiep": "IS.EXPENSE.GA",
        "chi phi ban hang": "IS.EXPENSE.SELLING",
    },
    "BALANCE_SHEET": {
        "tong tai san": "BS.ASSETS.TOTAL",
        "tai san ngan han": "BS.ASSETS.SHORT_TERM",
        "tai san dai han": "BS.ASSETS.LONG_TERM",
        "tien va tuong duong tien": "BS.ASSETS.CASH_AND_EQUIVALENTS",
        "tien va cac khoan tuong duong tien": "BS.ASSETS.CASH_AND_EQUIVALENTS",
        "cac khoan phai thu": "BS.ASSETS.RECEIVABLES",
        "cac khoan phai thu ngan han": "BS.ASSETS.RECEIVABLES",
        "hang ton kho rong": "BS.ASSETS.INVENTORY",
        "hang ton kho": "BS.ASSETS.INVENTORY",
        "von chu so huu": "BS.EQUITY.TOTAL",
        "tong von chu so huu": "BS.EQUITY.TOTAL",
        "no phai tra": "BS.LIABILITIES.TOTAL",
        "tong no phai tra": "BS.LIABILITIES.TOTAL",
        "no ngan han": "BS.LIABILITIES.SHORT_TERM",
        "no dai han": "BS.LIABILITIES.LONG_TERM",
        "vay va no thue tai chinh ngan han": "BS.LIABILITIES.SHORT_TERM_BORROWINGS",
        "vay ngan han": "BS.LIABILITIES.SHORT_TERM_BORROWINGS",
        "vay va no thue tai chinh dai han": "BS.LIABILITIES.LONG_TERM_BORROWINGS",
        "vay dai han": "BS.LIABILITIES.LONG_TERM_BORROWINGS",
    },
    "CASH_FLOW": {
        "luu chuyen tien thuan tu hoat dong kinh doanh": "CF.OPERATING.NET",
        "dong tien tu hoat dong kinh doanh": "CF.OPERATING.NET",
        "luu chuyen tien thuan tu hoat dong dau tu": "CF.INVESTING.NET",
        "luu chuyen tien thuan tu hoat dong tai chinh": "CF.FINANCING.NET",
        "tien chi mua sam tscd": "CF.CAPEX",
        "tien chi de mua sam xay dung tscd": "CF.CAPEX",
        "chi mua sam tai san co dinh": "CF.CAPEX",
        "khau hao tscd": "CF.OPERATING.DEPRECIATION",
        "khau hao tai san co dinh": "CF.OPERATING.DEPRECIATION",
    },
}


def normalize_string(val: Any) -> str:
    """Normalize string by converting Vietnamese accents and lowering case."""
    text = str(val or "").strip()
    text = text.lower()
    text = re.sub(r"[àáảãạăằắẳẵặâầấẩẫậ]", "a", text)
    text = re.sub(r"[èéẻẽẹêềếểễệ]", "e", text)
    text = re.sub(r"[ìíỉĩị]", "i", text)
    text = re.sub(r"[òóỏõọôồốổỗộơờớởỡợ]", "o", text)
    text = re.sub(r"[ùúủũụưừứửữự]", "u", text)
    text = re.sub(r"[ỳýỷỹỵ]", "y", text)
    text = re.sub(r"[đ]", "d", text)
    text = re.sub(r"[^a-z0-9\s]", "", text)
    return re.sub(r"\s+", " ", text).strip()


@dataclass
class SSIFileMetadata:
    filename: str
    symbol: str
    statement_type: str
    export_date_str: str | None
    download_suffix: int | None
    is_valid: bool


def parse_ssi_filename(filename: str) -> SSIFileMetadata:
    """Parse filename metadata tolerant of casing and download suffixes.

    Example: SSI_AAA_Financial_statement_Cash_Flow_09092026 (5).xlsx
    => symbol: AAA, statement_type: CASH_FLOW, export_date: 09092026, suffix: 5
    """
    basename = Path(filename).name
    # Match pattern: SSI_<SYMBOL>_Financial_[Ss]tatement_<STMT_TYPE>_<DATE>[ _]?(<SUFFIX>).xlsx
    match = re.search(
        r"^SSI_([A-Z0-9]{3,10})_Financial_[Ss]tatement_([A-Za-z_]+)_([0-9]{8})(?:\s*\(([0-9]+)\))?\.xlsx$",
        basename,
        re.IGNORECASE,
    )
    if not match:
        return SSIFileMetadata(
            filename=basename,
            symbol="UNKNOWN",
            statement_type="UNKNOWN",
            export_date_str=None,
            download_suffix=None,
            is_valid=False,
        )

    symbol = match.group(1).upper()
    raw_stmt = match.group(2).upper().replace(" ", "_")
    export_date_str = match.group(3)
    suffix = int(match.group(4)) if match.group(4) else None

    stmt_type = STATEMENT_TYPE_MAP.get(raw_stmt, "UNKNOWN")
    if stmt_type == "UNKNOWN":
        for k, v in STATEMENT_TYPE_MAP.items():
            if k in raw_stmt:
                stmt_type = v
                break

    return SSIFileMetadata(
        filename=basename,
        symbol=symbol,
        statement_type=stmt_type,
        export_date_str=export_date_str,
        download_suffix=suffix,
        is_valid=(stmt_type != "UNKNOWN"),
    )


@dataclass
class SSIParsedWorkbook:
    metadata: SSIFileMetadata
    file_path: str
    raw_file_sha256: str
    semantic_hash: str
    extraction_timestamp: str | None
    periods: list[str]
    observations: list[dict[str, Any]]  # list of {raw_line_name, period, value}
    is_valid: bool
    error: str | None = None


def parse_ssi_workbook(file_path: str) -> SSIParsedWorkbook:
    """Parse workbook content, extract financial values, compute semantic hash."""
    path = Path(file_path)
    meta = parse_ssi_filename(path.name)

    with open(path, "rb") as f:
        raw_bytes = f.read()
    raw_file_sha256 = hashlib.sha256(raw_bytes).hexdigest()

    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb.active
        sheet_rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()

        def get_val(r_1based: int, c_1based: int):
            r_idx = r_1based - 1
            c_idx = c_1based - 1
            if 0 <= r_idx < len(sheet_rows):
                row_data = sheet_rows[r_idx]
                if 0 <= c_idx < len(row_data):
                    return row_data[c_idx]
            return None

        # Extract timestamp in row 6 column 2 if present
        extract_ts = None
        cell_ts = get_val(6, 2)
        if isinstance(cell_ts, datetime):
            extract_ts = cell_ts.isoformat()
        elif cell_ts:
            extract_ts = str(cell_ts)

        # Locate header row containing periods (Row 8 usually)
        header_row_idx = 8
        periods: list[str] = []
        period_col_map: dict[int, str] = {}

        for col in range(2, 50):
            val = get_val(header_row_idx, col)
            if val is not None and str(val).strip():
                p_str = str(val).strip()
                periods.append(p_str)
                period_col_map[col] = p_str

        if not periods:
            # Fallback search for period row
            for r in range(1, 12):
                row_vals = [get_val(r, c) for c in range(2, 20)]
                year_like = [str(v).strip() for v in row_vals if v and re.match(r"^(19|20)\d{2}", str(v).strip())]
                if len(year_like) >= 2:
                    header_row_idx = r
                    periods = []
                    period_col_map = {}
                    for col in range(2, 50):
                        v = get_val(r, col)
                        if v is not None and str(v).strip():
                            p_str = str(v).strip()
                            periods.append(p_str)
                            period_col_map[col] = p_str
                    break

        if not periods:
            return SSIParsedWorkbook(
                metadata=meta,
                file_path=str(path),
                raw_file_sha256=raw_file_sha256,
                semantic_hash="",
                extraction_timestamp=extract_ts,
                periods=[],
                observations=[],
                is_valid=False,
                error="Could not locate period header row in workbook.",
            )

        observations: list[dict[str, Any]] = []

        # Read data rows starting after header row
        for row in range(header_row_idx + 1, len(sheet_rows) + 1):
            line_name = get_val(row, 1)
            if line_name is None or not str(line_name).strip():
                continue
            line_name_str = str(line_name).strip()

            for col, period_str in period_col_map.items():
                cell_val = get_val(row, col)
                # Preserve NULL vs 0. Empty cell or None is None.
                numeric_val = None
                if cell_val is not None and cell_val != "":
                    try:
                        numeric_val = float(cell_val)
                    except (ValueError, TypeError):
                        numeric_val = None

                observations.append({
                    "raw_line_name": line_name_str,
                    "period": period_str,
                    "value": numeric_val,
                })

        # Compute semantic payload hash (ignoring file metadata, extract timestamp, download suffix)
        canonical_payload = {
            "symbol": meta.symbol,
            "statement_type": meta.statement_type,
            "periods": sorted(periods),
            "observations": sorted(
                [
                    {
                        "raw_line_name": normalize_string(obs["raw_line_name"]),
                        "period": obs["period"],
                        "value": obs["value"],
                    }
                    for obs in observations
                    if obs["value"] is not None
                ],
                key=lambda x: (x["raw_line_name"], x["period"]),
            ),
        }
        semantic_hash = hashlib.sha256(json.dumps(canonical_payload, sort_keys=True).encode("utf-8")).hexdigest()

        return SSIParsedWorkbook(
            metadata=meta,
            file_path=str(path),
            raw_file_sha256=raw_file_sha256,
            semantic_hash=semantic_hash,
            extraction_timestamp=extract_ts,
            periods=periods,
            observations=observations,
            is_valid=True,
        )

    except Exception as exc:
        return SSIParsedWorkbook(
            metadata=meta,
            file_path=str(path),
            raw_file_sha256=raw_file_sha256,
            semantic_hash="",
            extraction_timestamp=None,
            periods=[],
            observations=[],
            is_valid=False,
            error=str(exc),
        )


def map_ssi_line_item(raw_line_name: str, statement_type: str) -> tuple[str | None, str]:
    """Map Vietnamese raw line name to canonical line item code.

    Returns (canonical_code or None, mapping_status: 'MAPPED' | 'UNMAPPED').
    """
    normalized = normalize_string(raw_line_name)
    mapping_dict = CANONICAL_LINE_MAPPINGS.get(statement_type, {})
    canonical = mapping_dict.get(normalized)
    if canonical:
        return canonical, "MAPPED"

    # Secondary fuzzy prefix/token match
    for norm_pattern, code in mapping_dict.items():
        if norm_pattern in normalized or normalized.startswith(norm_pattern):
            return code, "MAPPED"

    return None, "UNMAPPED"


def _parse_file_top(fp_str: str) -> SSIParsedWorkbook:
    return parse_ssi_workbook(fp_str)


@dataclass
class ImportSummaryReport:
    scanned_at: str
    source_directory: str
    files_discovered: int
    files_parsed: int
    parse_failures: int
    unique_symbols: int
    balance_sheet_payloads: int
    income_statement_payloads: int
    cash_flow_payloads: int
    same_payload_group_count: int
    different_payload_conflicts: int
    raw_observations_count: int
    canonical_facts_count: int
    unmapped_rows_count: int
    symbols: list[str] = field(default_factory=list)


class SSIBulkImporter:
    """Bulk SSI Ingestion Engine."""

    def __init__(self, directory: str, db_connection=None):
        self.directory = Path(directory)
        self.db = db_connection

    def run(
        self,
        dry_run: bool = True,
        resume: bool = True,
        target_symbol: str | None = None,
        batch_size: int = 200,
    ) -> ImportSummaryReport:
        if not self.directory.exists():
            raise FileNotFoundError(f"Directory {self.directory} does not exist.")

        if target_symbol:
            sym_clean = target_symbol.strip().upper()
            all_files = sorted([f for f in self.directory.glob(f"*{sym_clean}*.xlsx")])
        else:
            all_files = sorted([f for f in self.directory.glob("*.xlsx")])

        files_discovered = len(all_files)
        parsed_workbooks: list[SSIParsedWorkbook] = []
        parse_failures = 0
        symbols_found = set()

        semantic_groups: dict[str, list[SSIParsedWorkbook]] = {}
        processed_shas = set()

        if not dry_run and resume and self.db:
            existing_shas = self.db.execute("SELECT raw_file_sha256 FROM ssi_import_files WHERE import_status='SUCCESS'").fetchall()
            processed_shas = {r["raw_file_sha256"] for r in existing_shas}

        files_to_parse = [f for f in all_files if not f.name.startswith("~$")]
        print(f"Parsing {len(files_to_parse)} SSI XLSX workbooks using parallel workers...", flush=True)

        def _parse_file(fp: Path) -> SSIParsedWorkbook:
            return parse_ssi_workbook(str(fp))

        parsed_results = []
        with ThreadPoolExecutor(max_workers=min(16, (os.cpu_count() or 4) * 2)) as executor:
            futures = [executor.submit(_parse_file, f) for f in files_to_parse]
            for i, fut in enumerate(as_completed(futures), 1):
                parsed_results.append(fut.result())
                if i % 500 == 0 or i == len(files_to_parse):
                    print(f"[{i}/{len(files_to_parse)}] Excel workbooks parsed...", flush=True)

        for parsed in parsed_results:
            if not parsed.is_valid:
                parse_failures += 1
                parsed_workbooks.append(parsed)
                continue

            symbols_found.add(parsed.metadata.symbol)
            parsed_workbooks.append(parsed)
            semantic_groups.setdefault(parsed.semantic_hash, []).append(parsed)

        print(f"Parsed {len(parsed_workbooks)} workbooks across {len(symbols_found)} symbols.", flush=True)

        bs_count = sum(1 for p in parsed_workbooks if p.metadata.statement_type == "BALANCE_SHEET" and p.is_valid)
        is_count = sum(1 for p in parsed_workbooks if p.metadata.statement_type == "INCOME_STATEMENT" and p.is_valid)
        cf_count = sum(1 for p in parsed_workbooks if p.metadata.statement_type == "CASH_FLOW" and p.is_valid)

        same_payload_groups = sum(1 for g in semantic_groups.values() if len(g) > 1)

        raw_obs_total = 0
        canonical_facts_total = 0
        unmapped_total = 0

        # Dry run report generation
        if dry_run or not self.db:
            for p in parsed_workbooks:
                if not p.is_valid:
                    continue
                for obs in p.observations:
                    raw_obs_total += 1
                    code, status = map_ssi_line_item(obs["raw_line_name"], p.metadata.statement_type)
                    if status == "UNMAPPED":
                        unmapped_total += 1
                    else:
                        canonical_facts_total += 1

            summary = ImportSummaryReport(
                scanned_at=datetime.now(timezone.utc).isoformat(),
                source_directory=str(self.directory),
                files_discovered=files_discovered,
                files_parsed=len(parsed_workbooks) - parse_failures,
                parse_failures=parse_failures,
                unique_symbols=len(symbols_found),
                balance_sheet_payloads=bs_count,
                income_statement_payloads=is_count,
                cash_flow_payloads=cf_count,
                same_payload_group_count=same_payload_groups,
                different_payload_conflicts=0,
                raw_observations_count=raw_obs_total,
                canonical_facts_count=canonical_facts_total,
                unmapped_rows_count=unmapped_total,
                symbols=sorted(list(symbols_found)),
            )
            self.write_reports(summary, parsed_workbooks)
            return summary

        # Production Database Import Phase
        now_str = datetime.now(timezone.utc).isoformat()
        total_workbooks = len(parsed_workbooks)

        for idx, parsed in enumerate(parsed_workbooks, 1):
            if idx % 100 == 0 or idx == total_workbooks:
                print(f"[{idx}/{total_workbooks}] Processing SSI database import: symbol={parsed.metadata.symbol} statement={parsed.metadata.statement_type}")

            if not parsed.is_valid:
                self.db.execute(
                    """INSERT INTO ssi_import_files
                       (symbol, statement_type, file_name, raw_file_sha256, semantic_hash, import_status, imported_at, error)
                       VALUES (?, ?, ?, ?, ?, 'PARSE_ERROR', ?, ?)
                       ON CONFLICT (raw_file_sha256) DO NOTHING""",
                    (parsed.metadata.symbol, parsed.metadata.statement_type, parsed.metadata.filename, parsed.raw_file_sha256, "", now_str, parsed.error or "Parse error"),
                )
                continue

            if parsed.raw_file_sha256 in processed_shas:
                continue

            # Insert ssi_import_files record
            res = self.db.execute(
                """INSERT INTO ssi_import_files
                   (symbol, statement_type, file_name, raw_file_sha256, semantic_hash, export_date, extraction_timestamp, import_status, imported_at)
                   VALUES (?, ?, ?, ?, ?, ?, ?, 'SUCCESS', ?)
                   ON CONFLICT (raw_file_sha256) DO UPDATE SET import_status='SUCCESS' RETURNING id""",
                (parsed.metadata.symbol, parsed.metadata.statement_type, parsed.metadata.filename, parsed.raw_file_sha256, parsed.semantic_hash, parsed.metadata.export_date_str, parsed.extraction_timestamp, now_str),
            ).fetchone()
            file_id = res["id"] if res else None

            raw_obs_batch = []
            canonical_batch = []

            for obs in parsed.observations:
                raw_obs_total += 1
                line_name = obs["raw_line_name"]
                period = obs["period"]
                val = obs["value"]

                fiscal_year = int(period) if re.match(r"^\d{4}$", period) else None
                fiscal_quarter = None

                code, map_status = map_ssi_line_item(line_name, parsed.metadata.statement_type)
                if map_status == "UNMAPPED":
                    unmapped_total += 1
                else:
                    canonical_facts_total += 1

                if file_id:
                    raw_obs_batch.append((
                        file_id, parsed.metadata.symbol, parsed.metadata.statement_type, line_name, period, fiscal_year, fiscal_quarter, val, map_status, code, now_str
                    ))

                if code and val is not None and fiscal_year:
                    canonical_batch.append((
                        parsed.metadata.symbol, parsed.metadata.statement_type, code, val, fiscal_year, f"{fiscal_year}-12-31", now_str
                    ))

            if file_id and raw_obs_batch:
                self.db.executemany(
                    """INSERT INTO ssi_raw_financial_observations
                       (import_file_id, symbol, statement_type, raw_line_name, period, fiscal_year, fiscal_quarter, value, mapping_status, canonical_code, created_at)
                       VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                    raw_obs_batch,
                )

            if canonical_batch:
                sym = canonical_batch[0][0]
                st_type = canonical_batch[0][1]
                self.db.execute(
                    """DELETE FROM canonical_facts
                       WHERE symbol = ? AND statement_type = ? AND provider = 'ssi'""",
                    (sym, st_type),
                )
                self.db.executemany(
                    """INSERT INTO canonical_facts
                       (symbol, statement_type, line_item_code, value, period_type, fiscal_year, fiscal_quarter, period_end, provider, quality_status, observed_at)
                       VALUES (?, ?, ?, ?, 'FY', ?, NULL, ?, 'ssi', 'PRIMARY_SSI', ?)""",
                    canonical_batch,
                )

            if self.db and hasattr(self.db, "commit"):
                self.db.commit()

        if self.db and hasattr(self.db, "commit"):
            self.db.commit()

        summary = ImportSummaryReport(
            scanned_at=now_str,
            source_directory=str(self.directory),
            files_discovered=files_discovered,
            files_parsed=len(parsed_workbooks) - parse_failures,
            parse_failures=parse_failures,
            unique_symbols=len(symbols_found),
            balance_sheet_payloads=bs_count,
            income_statement_payloads=is_count,
            cash_flow_payloads=cf_count,
            same_payload_group_count=same_payload_groups,
            different_payload_conflicts=0,
            raw_observations_count=raw_obs_total,
            canonical_facts_count=canonical_facts_total,
            unmapped_rows_count=unmapped_total,
            symbols=sorted(list(symbols_found)),
        )
        self.write_reports(summary, parsed_workbooks)
        return summary

    def write_reports(self, summary: ImportSummaryReport, workbooks: list[SSIParsedWorkbook]) -> None:
        reports_dir = Path("docs/reports")
        reports_dir.mkdir(parents=True, exist_ok=True)

        with open(reports_dir / "ssi-import-summary.json", "w", encoding="utf-8") as f:
            json.dump(asdict(summary), f, indent=2, ensure_ascii=False)

        unmapped_set = set()
        errors = []

        for wb in workbooks:
            if not wb.is_valid and wb.error:
                errors.append({"file": wb.metadata.filename, "symbol": wb.metadata.symbol, "error": wb.error})
            for obs in wb.observations:
                code, status = map_ssi_line_item(obs["raw_line_name"], wb.metadata.statement_type)
                if status == "UNMAPPED":
                    unmapped_set.add((wb.metadata.statement_type, obs["raw_line_name"]))

        with open(reports_dir / "ssi-unmapped-line-items.csv", "w", encoding="utf-8", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(["statement_type", "raw_line_name"])
            for st, line in sorted(list(unmapped_set)):
                writer.writerow([st, line])

        with open(reports_dir / "ssi-import-errors.csv", "w", encoding="utf-8", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(["file", "symbol", "error"])
            for err in errors:
                writer.writerow([err["file"], err["symbol"], err["error"]])


def main():
    parser = argparse.ArgumentParser(description="Bulk SSI Financial Ingestion Importer")
    parser.add_argument("--directory", type=str, default="F:\\DATA_BCTC", help="Target source directory containing SSI XLSX files")
    parser.add_argument("--dry-run", action="store_true", help="Perform discovery and parsing without database mutation")
    parser.add_argument("--import", dest="do_import", action="store_true", help="Execute database import")
    parser.add_argument("--resume", action="store_true", help="Resume from previous database import state")
    parser.add_argument("--symbol", type=str, default=None, help="Filter import to specific ticker symbol")

    args = parser.parse_args()

    is_dry_run = args.dry_run or not args.do_import

    if not is_dry_run:
        from portfolio.finance_catalog import initialize_finance_schema
        initialize_finance_schema()
        with _schema_connection(FINANCE_SCHEMA) as conn:
            importer = SSIBulkImporter(directory=args.directory, db_connection=conn)
            summary = importer.run(dry_run=is_dry_run, resume=args.resume, target_symbol=args.symbol)
    else:
        importer = SSIBulkImporter(directory=args.directory, db_connection=None)
        summary = importer.run(dry_run=True, resume=args.resume, target_symbol=args.symbol)

    print("=== SSI BULK INGESTION SUMMARY ===")
    print(f"Directory: {summary.source_directory}")
    print(f"Files Discovered: {summary.files_discovered}")
    print(f"Files Parsed: {summary.files_parsed}")
    print(f"Parse Failures: {summary.parse_failures}")
    print(f"Unique Symbols: {summary.unique_symbols}")
    print(f"Balance Sheet Payloads: {summary.balance_sheet_payloads}")
    print(f"Income Statement Payloads: {summary.income_statement_payloads}")
    print(f"Cash Flow Payloads: {summary.cash_flow_payloads}")
    print(f"Same-Financial-Payload Groups: {summary.same_payload_group_count}")
    print(f"Raw Observations Count: {summary.raw_observations_count}")
    print(f"Canonical Facts Mapped: {summary.canonical_facts_count}")
    print(f"Unmapped Line Items Count: {summary.unmapped_rows_count}")


if __name__ == "__main__":
    main()
